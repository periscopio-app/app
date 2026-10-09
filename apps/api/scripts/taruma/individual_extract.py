#!/usr/bin/env python3
"""
Extrator LOCAL da base individual de Tarumã (planilha BD NEMT) para o contrato de importação do Periscópio.

Roda na máquina de quem tem a planilha e a autorização da controladora dos dados. O dado de criança real
não precisa passar por nenhum servidor além do banco de destino.

  1) sugerir o mapeamento (lê SÓ a linha de cabeçalho, nenhum dado):
       python3 -I individual_extract.py suggest planilha.xlsx > mapping.json
     Revise o mapping.json (campos "_todo") antes de seguir.

  2) extrair:
       PSEUDONYM_SECRET='<32+ caracteres, guardado fora do banco>' MAPBOX_TOKEN='<token>' \
       python3 -I individual_extract.py extract planilha.xlsx --mapping mapping.json --out individual.ndjson
     --geocode off   não chama o Mapbox (todas as localizações ficam "skipped")
     --no-permanent  não usa permanent=true (o termo do Mapbox só permite guardar o resultado com a geocodificação permanente)

O que acontece com cada dado:
  - nome da criança  -> UUID (HMAC-SHA256 com PSEUDONYM_SECRET). O nome nunca é gravado.
  - endereço         -> latitude/longitude (Mapbox Geocoding v6). O endereço nunca é gravado nem impresso.
  - nomes de outras pessoas (ex.: familiar responsável, quando é nome) -> UUID, o mesmo UUID para a mesma pessoa.
  - todo o resto     -> migrado como está, por coluna (camada bruta) e nos campos do contrato.
Guarde o PSEUDONYM_SECRET: é a única forma de religar um UUID a uma criança (recalculando o HMAC a partir do nome).
"""
import argparse
import collections
import datetime as dt
import hashlib
import hmac
import json
import os
import re
import sys
import time
import unicodedata
import urllib.error
import urllib.parse
import urllib.request
import uuid

SCHEMA_VERSION = 1
DEFAULT_MAPBOX = "https://api.mapbox.com"
# Tarumã (SP): centro aproximado do município e caixa de segurança (lon_min, lat_min, lon_max, lat_max)
DEFAULT_PROXIMITY = (-50.5811, -22.7467)
DEFAULT_BBOX = (-51.10, -23.20, -50.10, -22.30)

COMPLAINTS = ["Não devidamente esclarecida", "Falta de atenção", "Hiperatividade",
              "Hiperatividade + falta de atenção", "Escrita", "Leitura", "Escrita + leitura",
              "Indisciplina", "Agressividade", "Atraso de fala e linguagem", "Não compreende",
              "Dificuldade em matemática", "Transtorno de ajustamento", "Ansiedade com ou sem tiques",
              "Compulsões/obsessões com ou sem tiques", "Humor deprimido"]
SERVICES = {
    "fonoaudiologia": ["FONOID", "FONOaval", "FONOter"],
    "psicopedagogia": ["PpCaval", "PpCter"],
    "psicoterapia": ["PSICOter"],
    "psicomotricidade": ["PsicoM aval", "PsicoM ter"],
    "neuropsicologia": ["ANP"],
    "assistencia_social": ["AS"],
    "consulta_medica": ["SESSÕES MD 1", "SESSÕES MD 2", "SESSÕES MD 3"],
}
SCHOOLS = {"0": "Não se aplica", "1": "MAB", "2": "GL", "3": "JOO", "4": "JR", "5": "HH", "6": "RR",
           "7": "IOF", "8": "DONA COTA", "9": "VILA DO LAGO", "10": "DAVID LUZ", "11": "SJ"}


class ExtractError(Exception):
    pass


# ── utilitários ──────────────────────────────────────────────────────────────────────
def norm(s):
    s = unicodedata.normalize("NFKD", str(s))
    s = "".join(c for c in s if not unicodedata.combining(c)).lower()
    return re.sub(r"[^a-z0-9]+", " ", s).strip()


def sha256_file(path):
    h = hashlib.sha256()
    with open(path, "rb") as fh:
        for block in iter(lambda: fh.read(1 << 20), b""):
            h.update(block)
    return h.hexdigest()


def pseudo_uuid(secret, *parts):
    """UUID estável derivado por HMAC: mesma entrada + mesmo segredo = mesmo UUID. Não é reversível sem o nome."""
    msg = "\x1f".join(norm(p) for p in parts).encode("utf-8")
    raw = bytearray(hmac.new(secret.encode("utf-8"), msg, hashlib.sha256).digest()[:16])
    raw[6] = (raw[6] & 0x0F) | 0x50  # versão 5
    raw[8] = (raw[8] & 0x3F) | 0x80  # variante RFC 4122
    return str(uuid.UUID(bytes=bytes(raw)))


def hmac_hex(secret, label, value):
    return hmac.new(secret.encode("utf-8"), f"{label}\x1f{norm(value)}".encode("utf-8"), hashlib.sha256).hexdigest()


def empty(v):
    return v is None or (isinstance(v, str) and v.strip() == "")


def jsonable(v):
    if isinstance(v, dt.datetime):
        return v.date().isoformat() if (v.hour, v.minute, v.second) == (0, 0, 0) else v.isoformat()
    if isinstance(v, dt.date):
        return v.isoformat()
    if isinstance(v, float) and v.is_integer():
        return int(v)
    if isinstance(v, bool):
        return int(v)
    if isinstance(v, (int, float, str)) or v is None:
        return v.strip() if isinstance(v, str) else v
    return str(v)


def to_text(v, limit=None):
    if empty(v):
        return None
    t = str(jsonable(v)).strip()
    return t[:limit] if limit else t


def to_int(v):
    if empty(v):
        return None
    try:
        return int(float(str(v).replace(",", ".")))
    except ValueError:
        return None


DATE_FORMATS = ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y", "%d/%m/%y", "%d.%m.%Y")


def to_date(v):
    """Devolve (iso|None, ok). ok=False quando havia valor e não foi possível ler como data."""
    if empty(v):
        return None, True
    if isinstance(v, dt.datetime):
        return v.date().isoformat(), True
    if isinstance(v, dt.date):
        return v.isoformat(), True
    s = str(v).strip()
    for f in DATE_FORMATS:
        try:
            return dt.datetime.strptime(s, f).date().isoformat(), True
        except ValueError:
            continue
    return None, False


# ── planilha ─────────────────────────────────────────────────────────────────────────
def open_sheet(path, sheet):
    import openpyxl  # importado aqui para o --help funcionar sem a dependência
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if sheet not in wb.sheetnames:
        raise ExtractError(f"Aba '{sheet}' não existe. Abas: {', '.join(wb.sheetnames)}")
    return wb[sheet]


def unique_headers(row):
    seen, out = collections.Counter(), []
    for i, h in enumerate(row):
        name = str(h).strip() if h is not None and str(h).strip() else f"coluna_{i + 1}"
        seen[name] += 1
        out.append(name if seen[name] == 1 else f"{name}.{seen[name]}")
    return out


# ── sugestão de mapeamento (só cabeçalho) ───────────────────────────────────────────
def suggest(path, sheet):
    ws = open_sheet(path, sheet)
    header = unique_headers(next(ws.iter_rows(values_only=True)))
    used = set()

    def find(*patterns, exclude=()):
        for h in header:
            n = norm(h)
            if h not in used and any(re.search(p, n) for p in patterns) and not any(re.search(e, n) for e in exclude):
                return h
        return None

    def take(h):
        if h:
            used.add(h)
        return h

    name = take(find(r"^nome$", r"^nome (do )?(aluno|paciente|crianca)", r"^(aluno|paciente|crianca)$"))
    mapping = {
        "_comment": "Rascunho gerado só pelos cabeçalhos. Revise tudo; campos null = não encontrado (preencha ou deixe null).",
        "schemaVersion": SCHEMA_VERSION, "source": "taruma-nemt-2024", "referenceYear": 2024, "sheet": sheet,
        "rowFilterColumn": header[0],
        "school": {"column": take(find(r"^escola$")), "labels": SCHOOLS},
        "patient": {
            "nameColumn": name,
            "recordNumber": take(find(r"prontu")),
            "birthDate": take(find(r"nasc")),
            "currentAge": take(find(r"^id atual$", r"^idade")),
            "entryYear": take(find(r"^ano$")),
            "guardian": {"column": take(find(r"respons", r"familiar")), "as": "pseudonym", "relationColumn": take(find(r"parentesco", r"grau")),
                         "_todo": "'pseudonym' = coluna traz o NOME (vira UUID) e o parentesco sai de relationColumn, ou da própria célula se ela citar mãe/pai/avó...; 'text' = coluna traz só o parentesco"},
            "religion": take(find(r"relig")),
            "parentsOccupation": take(find(r"trabalho", r"profiss", r"ocupa")),
            "economicClass": take(find(r"classif.*econ", r"econom")),
            "extraColumns": [],
        },
        "address": {"columns": [h for h in (take(find(r"endere", r"logradouro")), take(find(r"bairro")), take(find(r"^cep$"))) if h],
                    "appendToQuery": "Tarumã, SP, Brasil",
                    "_todo": "appendToQuery é acrescentado à consulta; apague se os endereços já trazem a cidade ou se há moradores de outras cidades",
                    "bbox": list(DEFAULT_BBOX), "proximity": list(DEFAULT_PROXIMITY)},
        "complaints": {"columns": {f"QUEIXA-{i}": c for i, c in enumerate(COMPLAINTS) if f"QUEIXA-{i}" in header}, "flagValue": 1},
        "services": {k: [c for c in cols if c in header] for k, cols in SERVICES.items()},
        "items": {
            "hypothesis": {"column": take(find(r"hipotes", r"diagnost")), "split": r"\s*;\s*|\n"},
            "medication": {"column": take(find(r"farmac", r"medic")), "split": r"\s*;\s*|\n"},
            "family_history": {"column": take(find(r"anteced")), "split": r"\s*;\s*|\n"},
        },
        "dropColumns": [],
        "pseudonymColumns": [h for h in header if re.search(r"\bnome\b", norm(h)) and h != name and h not in used],
    }
    for c in mapping["complaints"]["columns"]:
        used.add(c)
    for cols in mapping["services"].values():
        used.update(cols)
    mapping["_unmappedColumns"] = [h for h in header if h not in used]
    return mapping


# ── geocodificação (Mapbox Geocoding v6) ────────────────────────────────────────────
class Geocoder:
    def __init__(self, token, base=None, permanent=True, bbox=DEFAULT_BBOX, proximity=DEFAULT_PROXIMITY, pause=0.06, suffix=""):
        self.token, self.base, self.suffix = token, (base or DEFAULT_MAPBOX).rstrip("/"), (suffix or "").strip()
        self.permanent, self.bbox, self.proximity, self.pause = permanent, bbox, proximity, pause
        self.cache, self.calls = {}, 0

    def lookup(self, address, key):
        if key in self.cache:
            return self.cache[key]
        res = self._call(f"{address}, {self.suffix}" if self.suffix else address)
        self.cache[key] = res
        return res

    def _call(self, address):
        q = {"q": address, "country": "br", "limit": "1", "language": "pt",
             "proximity": f"{self.proximity[0]},{self.proximity[1]}", "access_token": self.token}
        if self.permanent:
            q["permanent"] = "true"
        url = f"{self.base}/search/geocode/v6/forward?{urllib.parse.urlencode(q)}"
        body = None
        for attempt in range(6):
            time.sleep(self.pause)
            self.calls += 1
            try:
                with urllib.request.urlopen(urllib.request.Request(url), timeout=20) as r:
                    body = json.loads(r.read().decode("utf-8"))
                break
            except urllib.error.HTTPError as e:  # nunca imprime URL nem corpo (têm endereço e token)
                if e.code in (401, 403):
                    raise ExtractError(
                        f"Mapbox recusou a consulta (HTTP {e.code}). Confira o token; se usou permanent=true, a conta precisa "
                        "ter geocodificação permanente habilitada (ou rode com --no-permanent ciente do termo de uso).")
                if e.code == 429 or e.code >= 500:
                    time.sleep(min(30, 2 ** attempt))
                    continue
                return self._result("not_found")
            except (urllib.error.URLError, TimeoutError):
                time.sleep(min(30, 2 ** attempt))
        if body is None:
            raise ExtractError("Mapbox indisponível após várias tentativas; nada foi gravado. Tente de novo.")
        feats = body.get("features") or []
        if not feats:
            return self._result("not_found")
        f = feats[0]
        props = f.get("properties") or {}
        c = props.get("coordinates") or {}
        lat, lon = c.get("latitude"), c.get("longitude")
        if lat is None or lon is None:
            coords = (f.get("geometry") or {}).get("coordinates") or [None, None]
            lon, lat = coords[0], coords[1]
        if lat is None or lon is None:
            return self._result("not_found")
        accuracy = c.get("accuracy")
        confidence = (props.get("match_code") or {}).get("confidence")
        if not (self.bbox[0] <= lon <= self.bbox[2] and self.bbox[1] <= lat <= self.bbox[3]):
            return self._result("out_of_area", accuracy=accuracy, confidence=confidence)
        precise = accuracy in ("rooftop", "parcel", "point", "interpolated", "street") or confidence in ("exact", "high")
        return self._result("ok" if precise else "low_confidence", lat, lon, accuracy, confidence)

    def _result(self, status, lat=None, lon=None, accuracy=None, confidence=None):
        return {"status": status, "latitude": lat, "longitude": lon, "accuracy": accuracy, "confidence": confidence,
                "provider": "mapbox-geocoding-v6", "permanent": bool(self.permanent)}


# ── extração ─────────────────────────────────────────────────────────────────────────

KINSHIP = [  # (regex sobre texto sem acento e minúsculo, rótulo gravado)
    (r"\bmadrasta\b", "madrasta"), (r"\bpadrasto\b", "padrasto"),
    (r"\b(mae|mamae|mãe)\b", "mãe"), (r"\b(pai|papai)\b", "pai"),
    (r"\b(avo|vovo)\s*(materna|paterna)?\b", "avó/avô"), (r"\b(tia|tio)\b", "tia/tio"),
    (r"\b(irma|irmao)\b", "irmã/irmão"), (r"\b(tutor|tutora|responsavel legal|guardiao|guardia)\b", "responsável legal"),
]


def split_guardian(value, as_mode, relation_value=None):
    """Devolve (ref_uuid_base | None, parentesco | None). O NOME nunca sai daqui; só o texto-base para o HMAC."""
    text = to_text(value)
    rel = to_text(relation_value, 60)
    if not text:
        return None, rel
    if as_mode == "text":
        return None, (rel or text[:60])
    plain = unicodedata.normalize("NFD", text).encode("ascii", "ignore").decode().lower()
    found = None
    name_part = plain
    for rx, label in KINSHIP:
        if re.search(rx, plain):
            found = found or label
            name_part = re.sub(rx, " ", name_part)
    name_part = re.sub(r"[^a-z0-9]+", " ", name_part).strip()
    # célula só com parentesco ("mãe"): não há nome para pseudonimizar; vira parentesco
    if found and not name_part:
        return None, (rel or found)
    return (name_part or plain), (rel or found)


def build_manifest_and_lines(path, mapping_path, geocoder, secret):
    with open(mapping_path, encoding="utf-8") as fh:
        m = json.load(fh)
    mapping_sha = sha256_file(mapping_path)
    if m.get("schemaVersion") != SCHEMA_VERSION:
        raise ExtractError(f"schemaVersion do mapeamento deve ser {SCHEMA_VERSION}")
    ws = open_sheet(path, m.get("sheet", "Dados"))
    rows = list(ws.iter_rows(values_only=True))
    header = unique_headers(rows[0])
    ix = {h: i for i, h in enumerate(header)}

    pat = m["patient"]
    addr_cols = (m.get("address") or {}).get("columns") or []
    guardian = pat.get("guardian") or {}
    drop = list(m.get("dropColumns") or [])
    pseudo_cols = list(m.get("pseudonymColumns") or [])
    name_col = pat.get("nameColumn")
    if not name_col:
        raise ExtractError("O mapeamento precisa de patient.nameColumn (a coluna do nome da criança).")

    # todas as colunas citadas precisam existir no cabeçalho (só nomes de coluna aparecem na mensagem)
    cited = [name_col, (m.get("school") or {}).get("column"), m.get("rowFilterColumn"), *addr_cols, *drop, *pseudo_cols,
             pat.get("recordNumber"), pat.get("birthDate"), pat.get("currentAge"), pat.get("entryYear"),
             guardian.get("column"), guardian.get("relationColumn"), pat.get("religion"), pat.get("parentsOccupation"), pat.get("economicClass"),
             *(pat.get("extraColumns") or []), *list((m.get("complaints") or {}).get("columns", {}).keys()),
             *[c for cols in (m.get("services") or {}).values() for c in cols],
             *[(v or {}).get("column") for v in (m.get("items") or {}).values()]]
    missing = sorted({c for c in cited if c and c not in ix})
    if missing:
        raise ExtractError("Colunas do mapeamento que não existem na planilha: " + "; ".join(missing))
    school_col = (m.get("school") or {}).get("column")
    if not school_col:
        raise ExtractError("O mapeamento precisa de school.column.")

    filt = ix[m.get("rowFilterColumn") or header[0]]
    data = [(n, r) for n, r in enumerate(rows[1:], start=2) if r[filt] is not None and not empty(r[filt])]

    def cell(r, col):
        return r[ix[col]] if col else None

    transformed = [name_col, *addr_cols, *drop, *pseudo_cols] + ([guardian["column"]] if guardian.get("column") and guardian.get("as") == "pseudonym" else [])
    skip_raw = {name_col, *addr_cols, *drop}
    warnings, bad_dates = [], 0
    seen_ids = collections.Counter()
    labels = (m.get("school") or {}).get("labels") or {}
    comp_cols = (m.get("complaints") or {}).get("columns") or {}
    flag = (m.get("complaints") or {}).get("flagValue", 1)
    items_map = m.get("items") or {}

    nonnull = {h: 0 for h in header}
    lines = []
    for n, r in data:
        for h in header:
            if not empty(r[ix[h]]):
                nonnull[h] += 1
        name = to_text(cell(r, name_col))
        if not name:
            raise ExtractError(f"Linha {n} da planilha sem nome na coluna '{name_col}': não dá para gerar o UUID. Corrija a planilha.")
        birth, ok = to_date(cell(r, pat.get("birthDate")))
        if not ok:
            bad_dates += 1
        rec = to_text(cell(r, pat.get("recordNumber")), 60)
        base = pseudo_uuid(secret, "patient", name, birth or "", rec or "")
        seen_ids[base] += 1
        pid = base if seen_ids[base] == 1 else pseudo_uuid(secret, "patient", name, birth or "", rec or "", f"dup{seen_ids[base]}")

        sc = to_int(cell(r, school_col))
        school_code = str(sc if sc is not None else 0)
        patient = {
            "recordNumber": rec, "birthDate": birth,
            "currentAge": to_int(cell(r, pat.get("currentAge"))), "entryYear": to_int(cell(r, pat.get("entryYear"))),
            "religion": to_text(cell(r, pat.get("religion")), 80),
            "parentsOccupation": to_text(cell(r, pat.get("parentsOccupation")), 2000),
            "economicClass": to_text(cell(r, pat.get("economicClass")), 60),
        }
        g_base, g_rel = split_guardian(cell(r, guardian.get("column")), guardian.get("as"), cell(r, guardian.get("relationColumn")))
        if g_base:
            patient["guardianRef"] = pseudo_uuid(secret, "guardian", g_base)
        if g_rel:
            patient["guardianRelation"] = g_rel
        if pat.get("extraColumns"):
            patient["extra"] = {c: jsonable(cell(r, c)) for c in pat["extraColumns"] if not empty(cell(r, c))} or None

        # endereço -> coordenada (o texto nunca sai desta função)
        parts = [to_text(cell(r, c)) for c in addr_cols]
        address = ", ".join(p for p in parts if p)
        if not address:
            location = {"status": "no_address", "latitude": None, "longitude": None, "accuracy": None, "confidence": None,
                        "addressKey": None, "provider": None, "permanent": None}
        elif geocoder is None:
            location = {"status": "skipped", "latitude": None, "longitude": None, "accuracy": None, "confidence": None,
                        "addressKey": hmac_hex(secret, "address", address), "provider": None, "permanent": None}
        else:
            key = hmac_hex(secret, "address", address)
            location = {**geocoder.lookup(address, key), "addressKey": key}

        complaints = [label for col, label in comp_cols.items() if cell(r, col) == flag]
        services = []
        for svc, cols in (m.get("services") or {}).items():
            for col in cols:
                v = cell(r, col)
                if empty(v):
                    continue
                if isinstance(v, (int, float)) and not isinstance(v, bool):
                    services.append({"service": svc, "column": col, "valueNum": float(v), "valueText": None})
                elif isinstance(v, bool):
                    services.append({"service": svc, "column": col, "valueNum": float(int(v)), "valueText": None})
                else:
                    services.append({"service": svc, "column": col, "valueNum": None, "valueText": to_text(v)})
        items = []
        for kind, spec in items_map.items():
            text = to_text(cell(r, (spec or {}).get("column")))
            if not text:
                continue
            split = (spec or {}).get("split")
            for pos, piece in enumerate([p for p in (re.split(split, text) if split else [text]) if p and p.strip()]):
                items.append({"kind": kind, "position": pos, "text": piece.strip()})

        raw = {}
        for h in header:
            if h in skip_raw:
                continue
            v = r[ix[h]]
            if h == guardian.get("column") and guardian.get("as") == "pseudonym":
                gb, _ = split_guardian(v, "pseudonym")
                v = pseudo_uuid(secret, "guardian", gb) if gb else (patient.get("guardianRelation") if not empty(v) else None)
            elif h in pseudo_cols:
                v = pseudo_uuid(secret, "person", to_text(v)) if not empty(v) else None
            raw[h] = jsonable(v)

        lines.append({"patientId": pid, "sourceRow": n,
                      "school": {"code": school_code, "label": labels.get(school_code)},
                      "patient": patient, "location": location, "complaints": complaints,
                      "services": services, "items": items, "raw": raw})

    if bad_dates:
        warnings.append(f"{bad_dates} datas de nascimento não puderam ser lidas (ficaram nulas nos campos; o valor original segue na camada bruta)")
    dups = sum(1 for c in seen_ids.values() if c > 1)
    if dups:
        warnings.append(f"{dups} registros com nome, nascimento e prontuário idênticos receberam UUID distinto (possível cadastro duplicado)")
    for field in ("religion", "economicClass"):
        vals = [l["patient"].get(field) for l in lines if l["patient"].get(field)]
        if len(vals) >= 20 and len(set(vals)) / len(vals) > 0.6:
            warnings.append(f"o campo {field} tem muitos valores distintos ({len(set(vals))} de {len(vals)}): confira se a coluna não contém texto livre identificável")
    if geocoder is not None:
        st = collections.Counter(l["location"]["status"] for l in lines)
        warnings.append("localização: " + ", ".join(f"{k}={v}" for k, v in sorted(st.items())))

    manifest = {"manifest": {
        "schemaVersion": SCHEMA_VERSION, "source": m.get("source", "taruma-nemt-2024"), "referenceYear": m.get("referenceYear", 2024),
        "fileSha256": sha256_file(path), "mappingSha256": mapping_sha, "rowCount": len(lines),
        "columns": header, "transformedColumns": sorted(set(transformed)), "columnNonNull": nonnull,
        "geocoder": "mapbox-geocoding-v6" if geocoder else None, "warnings": warnings}}
    return manifest, lines


def write_ndjson(out, manifest, lines):
    old = os.umask(0o077)  # arquivo legível só pelo dono
    try:
        with open(out, "w", encoding="utf-8") as fh:
            fh.write(json.dumps(manifest, ensure_ascii=False) + "\n")
            for l in lines:
                fh.write(json.dumps(l, ensure_ascii=False) + "\n")
    finally:
        os.umask(old)


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)
    s = sub.add_parser("suggest", help="lê só o cabeçalho e imprime um mapeamento rascunho")
    s.add_argument("xlsx"); s.add_argument("--sheet", default="Dados")
    e = sub.add_parser("extract", help="gera o NDJSON do contrato")
    e.add_argument("xlsx"); e.add_argument("--mapping", required=True); e.add_argument("--out", required=True)
    e.add_argument("--geocode", choices=["mapbox", "off"], default="mapbox")
    e.add_argument("--no-permanent", action="store_true")
    a = ap.parse_args(argv)
    try:
        if a.cmd == "suggest":
            print(json.dumps(suggest(a.xlsx, a.sheet), ensure_ascii=False, indent=2))
            return 0
        secret = os.environ.get("PSEUDONYM_SECRET", "")
        if len(secret) < 32:
            raise ExtractError("Defina PSEUDONYM_SECRET com 32 ou mais caracteres (guarde-o fora do banco: é a chave de religação).")
        geocoder = None
        if a.geocode == "mapbox":
            token = os.environ.get("MAPBOX_TOKEN")
            if not token:
                raise ExtractError("Defina MAPBOX_TOKEN (ou use --geocode off).")
            with open(a.mapping, encoding="utf-8") as fh:
                ad = (json.load(fh).get("address") or {})
            geocoder = Geocoder(token, os.environ.get("MAPBOX_API_BASE"), permanent=not a.no_permanent,
                                bbox=tuple(ad.get("bbox") or DEFAULT_BBOX), proximity=tuple(ad.get("proximity") or DEFAULT_PROXIMITY),
                                suffix=ad.get("appendToQuery") or "")
        manifest, lines = build_manifest_and_lines(a.xlsx, a.mapping, geocoder, secret)
        write_ndjson(a.out, manifest, lines)
        print(f"{len(lines)} registros -> {a.out}. Nome e endereço não foram gravados.")
        for w in manifest["manifest"]["warnings"]:
            print("aviso:", w)
        return 0
    except ExtractError as err:
        print(f"ERRO: {err}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    sys.exit(main())
